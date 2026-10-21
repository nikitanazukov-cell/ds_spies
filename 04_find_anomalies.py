from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
from difflib import SequenceMatcher
import gzip
from itertools import groupby
import json
import math
from pathlib import Path
import re
import unicodedata
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

try:
    from unidecode import unidecode as _unidecode
    TRANSLITERATOR = 'Unidecode'
except ImportError:
    _unidecode = None
    TRANSLITERATOR = 'built-in Russian fallback'

_RU = dict(zip('абвгдеёжзийклмнопрстуфхцчшщъыьэюя',
    ['a','b','v','g','d','e','e','zh','z','i','i','k','l','m','n','o',
     'p','r','s','t','u','f','kh','ts','ch','sh','shch','','y','','e','yu','ya']))
_RU.update({k.upper():v.upper() for k,v in list(_RU.items())})

WEIGHTS = {
    'name_conflict': 6,
    'birthdate_conflict': 8,
    'flight_overlap': 10,
    'impossible_speed': 8,
    'short_airport_change': 5,
    'flight_route_conflict': 6,
    'over_20_flights_year': 1,
    'foreign_flight': 3,
}
CATEGORIES = {
    'Ложные документы': ('name_conflict', 'birthdate_conflict'),
    'Ложная запись о полёте': (
        'flight_overlap', 'impossible_speed', 'short_airport_change',
        'flight_route_conflict'),
    'Нетипичное поведение': ('over_20_flights_year', 'foreign_flight'),
}
DESCRIPTIONS = {
    'name_conflict': 'Один паспорт и существенно разные ФИО',
    'birthdate_conflict': 'Один паспорт и разные даты рождения',
    'flight_overlap': 'Пересекающиеся полёты или одновременный вылет в разные города',
    'impossible_speed': 'Нереальная скорость между аэропортами вылета',
    'short_airport_change': 'Нереально короткая смена аэропорта',
    'flight_route_conflict': 'Один номер рейса и дата, но разные точки маршрута',
    'over_20_flights_year': 'Более 20 рейсов за один календарный год',
    'foreign_flight': 'Есть международный перелёт',
}
COUNTRIES = {
    'RUSSIA':'RU','RUSSIAN FEDERATION':'RU','РОССИЯ':'RU','РФ':'RU','RUS':'RU',
    'UNITED KINGDOM':'GB','UK':'GB','GREAT BRITAIN':'GB','GBR':'GB',
    'UNITED STATES':'US','UNITED STATES OF AMERICA':'US','USA':'US',
    'FRANCE':'FR','FRA':'FR','GERMANY':'DE','DEU':'DE','ITALY':'IT','ITA':'IT','SPAIN':'ES','ESP':'ES',
    'CHINA':'CN','CHN':'CN','JAPAN':'JP','JPN':'JP','KAZAKHSTAN':'KZ','KAZ':'KZ','BELARUS':'BY',
    'TURKEY':'TR','TURKIYE':'TR','UAE':'AE','UNITED ARAB EMIRATES':'AE',
    'POLAND':'PL','UKRAINE':'UA','SWITZERLAND':'CH','AUSTRIA':'AT',
    'NETHERLANDS':'NL','BELGIUM':'BE','CZECH REPUBLIC':'CZ',
    'FINLAND':'FI','SWEDEN':'SE','NORWAY':'NO','DENMARK':'DK',
}
BASE_FIELDS = [
    'rank','person_id','passenger_name','passport','score_10','raw_points',
    'is_suspect','categories_ru','false_documents_points',
    'false_flight_points','unusual_behavior_points','score_breakdown','reasons_ru',
    'flights','max_flights_in_year','foreign_flights','flight_intervals_utc',
    *WEIGHTS.keys(),
]
EVENT_FIELDS = [
    'person_id','passenger_name','category','indicator','points','details',
    'flight_a','date_a','departure_a','arrival_a','origin_a','destination_a',
    'flight_b','date_b','departure_b','arrival_b','origin_b','destination_b',
    'source_a','source_record_a','source_b','source_record_b',
]


def open_csv(p, mode='r'):
    if str(p).lower().endswith('.gz'):
        return gzip.open(p,mode+'t',encoding='utf-8-sig',newline='')
    return open(p,mode,encoding='utf-8-sig',newline='')


def csv_separator(f):
    pos=f.tell()
    first=f.readline()
    f.seek(pos)
    return ';' if first.count(';') >= first.count(',') else ','


def transliterate(text):
    text=unicodedata.normalize('NFKC',(text or '').strip())
    if _unidecode:
        return _unidecode(text)
    return ''.join(_RU.get(c,c) for c in text)


def tokens_name(raw):
    s=transliterate(raw).upper().replace('Ё','Е')
    s=re.sub(r'[^A-Z]+',' ',s)
    words=[]
    for w in s.split():
        w=w.replace('X','KS').replace('YO','E').replace('J','I')
        if w in {'MR','MRS','MS','DR','G','JR','SR'}:
            continue
        if len(w)>=2:words.append(w)
    return tuple(sorted(words))


def same_name(a,b):
    wa,wb=tokens_name(a),tokens_name(b)
    if wa==wb:return True
    if len(wa)<2 or len(wb)<2:return True
    used=set();matches=0
    for x in sorted(wa,key=len,reverse=True):
        for j,y in enumerate(wb):
            if j in used:continue
            similarity=SequenceMatcher(None,x,y).ratio()
            if x==y or similarity>=0.82 or (min(len(x),len(y))>=4 and (x.startswith(y) or y.startswith(x))):
                matches+=1;used.add(j);break
    if matches>=2:return True
    return False


def norm_passport(raw):
    return re.sub(r'[^0-9A-Z]','',transliterate(raw).upper())


def country(code):
    c=re.sub(r'\s+',' ',(code or '').strip().upper())
    return COUNTRIES.get(c,c)


def airport_metadata(path):
    data={}
    if path:
        with open_csv(path) as f:
            rd=csv.DictReader(f,delimiter=csv_separator(f))
            if not {'iata','country','timezone'}.issubset(rd.fieldnames or []):
                raise ValueError('airport CSV must have iata;country;timezone (lat,lon,city optional)')
            for r in rd:
                code=(r.get('iata') or '').strip().upper()
                if code:
                    data[code]={'country':country(r.get('country')),
                        'tz':r.get('timezone') or '',
                        'lat':r.get('lat') or '', 'lon':r.get('lon') or '',
                        'city':r.get('city') or ''}
        return data,'--airport-meta CSV'
    try:
        import airportsdata
        for code,r in airportsdata.load('IATA').items():
            if len(code)==3:
                data[code.upper()]={'country':country(r.get('country')),
                    'tz':r.get('tz') or '', 'lat':r.get('lat'),
                    'lon':r.get('lon'), 'city':r.get('city') or ''}
        return data,'airportsdata'
    except ImportError:
        return {},'none (install airportsdata or supply --airport-meta)'


def utc_time(day, clock, airport_code, airport_info):
    if not(day and clock and airport_code):return None
    tz=airport_info.get(airport_code.upper(),{}).get('tz')
    if not tz:return None
    try:
        local=datetime.fromisoformat(day+'T'+clock)
        zone=ZoneInfo(tz)
        t0=local.replace(tzinfo=zone,fold=0)
        t1=local.replace(tzinfo=zone,fold=1)
        if t0.utcoffset()!=t1.utcoffset():return None
        t=t0.astimezone(timezone.utc)
        if t.astimezone(zone).replace(tzinfo=None)!=local:return None
        return t
    except (ValueError,ZoneInfoNotFoundError,OverflowError):
        return None


def haversine_km(a,b,airports):
    try:
        x,y=airports[a],airports[b]
        p1,l1,p2,l2=[math.radians(float(z)) for z in
            (x['lat'],x['lon'],y['lat'],y['lon'])]
        sin2=math.sin((p2-p1)/2)**2+math.cos(p1)*math.cos(p2)*math.sin((l2-l1)/2)**2
        return 2*6371*math.asin(min(1,math.sqrt(sin2)))
    except (KeyError,ValueError,TypeError):return None


def best_record(group):
    return max(group,key=lambda r:(sum(bool(r.get(k)) for k in
            ('arrival_date','arrival_time','departure_time','origin_airport',
             'destination_airport','origin_city','destination_city')),
        int(r.get('observation_count') or 1)))


def destination_city(r,airports):
    code=(r.get('destination_airport') or '').upper()
    s=(r.get('destination_city') or airports.get(code,{}).get('city') or '').strip()
    return re.sub(r'[^A-Z]','',transliterate(s).upper())


def different_destination_cities(a,b,airports):
    ca,cb=destination_city(a,airports),destination_city(b,airports)
    if ca and cb:
        return ca!=cb and SequenceMatcher(None,ca,cb).ratio()<0.87
    return False


def routes_disagree(a,b):
    for airport,city,country_col in (
        ('origin_airport','origin_city','origin_country'),
        ('destination_airport','destination_city','destination_country')):
        x,y=(a.get(airport) or '').upper(),(b.get(airport) or '').upper()
        if x and y and x!=y:return True
        cx,cy=transliterate(a.get(city) or '').upper(),transliterate(b.get(city) or '').upper()
        if cx and cy and cx!=cy and SequenceMatcher(None,cx,cy).ratio()<0.84:return True
        cxa,cya=country(a.get(country_col)),country(b.get(country_col))
        if cxa and cya and cxa!=cya:return True
    return False


def physical_key(r):
    base=(r.get('departure_date') or '',r.get('departure_time') or '',
          r.get('origin_airport') or '',r.get('destination_airport') or '')
    if all(base):return ('ITINERARY',)+base
    return ('NUMBER',r.get('departure_date') or '',r.get('flight_number') or '')


def event_row(pid,name,category,indicator,details,a=None,b=None):
    a=a or {};b=b or {}
    return {'person_id':pid,'passenger_name':name,'category':category,
        'indicator':indicator,'points':WEIGHTS[indicator], 'details':details,
        'flight_a':a.get('flight_number',''),'date_a':a.get('departure_date',''),
        'departure_a':a.get('departure_time',''),'arrival_a':a.get('arrival_time',''),
        'origin_a':a.get('origin_airport',''),'destination_a':a.get('destination_airport',''),
        'flight_b':b.get('flight_number',''),'date_b':b.get('departure_date',''),
        'departure_b':b.get('departure_time',''),'arrival_b':b.get('arrival_time',''),
        'origin_b':b.get('origin_airport',''),'destination_b':b.get('destination_airport',''),
        'source_a':a.get('source',''),'source_record_a':a.get('source_record_id',''),
        'source_b':b.get('source',''),'source_record_b':b.get('source_record_id','')}


def analyze_person(pid,observations,airports,args,write_event,stats):
    if not observations:return None
    display_name=next((r['passenger_name'] for r in observations if r['passenger_name']),'')
    docs={norm_passport(p) for r in observations for p in (r['passport'] or '').split('|') if p.strip()}
    if pid.startswith('P:'):
        docs.add(norm_passport(pid[2:]))
    docs.discard('')
    count=Counter()
    category_map={k:label for label,ind in CATEGORIES.items() for k in ind}

    def mark(indicator,details,a=None,b=None):
        count[indicator]+=1
        write_event(event_row(pid,display_name,category_map[indicator],indicator,details,a,b))

    by_passport=defaultdict(list)
    for r in observations:
        row_docs={norm_passport(p) for p in (r['passport'] or '').split('|') if p.strip()}
        if not row_docs and pid.startswith('P:'):
            row_docs={norm_passport(pid[2:])}
        for d in row_docs:
            if d:by_passport[d].append(r)

    for doc,related in by_passport.items():
        dates=sorted({r['date_of_birth'].strip() for r in related if r['date_of_birth'].strip()})
        if len(dates)>1:
            first=next(r for r in related if r['date_of_birth'].strip()==dates[0])
            second=next(r for r in related if r['date_of_birth'].strip()==dates[1])
            mark('birthdate_conflict',f'Passport {doc}: dates of birth {dates}',first,second)
        unique_names=[]
        for r in related:
            n=(r['passenger_name'] or '').strip()
            if n and len(tokens_name(n))>=2 and all(tokens_name(n)!=tokens_name(x[0]) for x in unique_names):
                unique_names.append((n,r))
        made=0
        for i,(n1,r1) in enumerate(unique_names):
            for n2,r2 in unique_names[i+1:]:
                if not same_name(n1,n2):
                    mark('name_conflict',f'Passport {doc}: "{n1}" vs "{n2}"',r1,r2)
                    made+=1
                    if made>=4:break
            if made>=4:break

    by_number_date=defaultdict(list)
    for r in observations:
        by_number_date[(r['departure_date'],r['flight_number'])].append(r)
    flights=[]
    for (day,number),variants in by_number_date.items():
        for i,r in enumerate(variants):
            conflicting=next((x for x in variants[i+1:] if routes_disagree(r,x)),None)
            if conflicting is not None:
                mark('flight_route_conflict',
                    f'{number}, {day}: different origin or destination',r,conflicting)
                break
        flights.append(best_record(variants).copy())
    physical={}
    for r in flights:
        k=physical_key(r)
        if k not in physical or best_record([physical[k],r]) is r:
            physical[k]=r
    flights=list(physical.values())
    years=Counter();foreign=0;country_comparable=0;full_intervals=0
    for r in flights:
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}',r['departure_date'] or ''):
            years[r['departure_date'][:4]]+=1
        aa=(r['origin_airport'] or '').upper();bb=(r['destination_airport'] or '').upper()
        oc=country(r['origin_country'] or airports.get(aa,{}).get('country',''))
        dc=country(r['destination_country'] or airports.get(bb,{}).get('country',''))
        if oc and dc:
            country_comparable+=1
            if oc!=dc:foreign+=1
        r['_dep_utc']=utc_time(r['departure_date'],r['departure_time'],aa,airports)
        r['_arr_utc']=utc_time(r['arrival_date'],r['arrival_time'],bb,airports)
        if r['_dep_utc'] and r['_arr_utc']:
            if r['_arr_utc']>r['_dep_utc']:
                full_intervals+=1
            else:r['_arr_utc']=None

    if years and max(years.values())>20:
        mark('over_20_flights_year',f'{max(years.values())} distinct flights in {max(years,key=years.get)}')
    if foreign:
        mark('foreign_flight',f'{foreign} cross-border flights out of {country_comparable} with known countries')

    departures=defaultdict(list)
    for r in flights:
        if r['departure_date'] and r['departure_time'] and r['origin_airport']:
            departures[(r['departure_date'],r['departure_time'],r['origin_airport'])].append(r)
    overlap_pairs=set()
    for group in departures.values():
        for i,a in enumerate(group):
            for b in group[i+1:]:
                if (a['flight_number']!=b['flight_number'] and
                        different_destination_cities(a,b,airports)):
                    ids=tuple(sorted((id(a),id(b))))
                    if ids not in overlap_pairs:
                        overlap_pairs.add(ids)
                        mark('flight_overlap','Simultaneous departures to distinct cities; not same-route codeshare',a,b)
                    break

    timed=sorted((r for r in flights if r['_dep_utc']),key=lambda x:x['_dep_utc'])
    active=[]
    same_utc_departures=[]
    for i,r in enumerate(timed):
        same_utc_departures=[old for old in same_utc_departures
            if (r['_dep_utc']-old['_dep_utc']).total_seconds()<60]
        for a in same_utc_departures:
            if a['flight_number']==r['flight_number'] and a['departure_date']==r['departure_date']:
                continue
            if not different_destination_cities(a,r,airports):continue
            ids=tuple(sorted((id(a),id(r))))
            if ids not in overlap_pairs:
                overlap_pairs.add(ids)
                mark('flight_overlap',
                    'Same UTC departure instant to different cities; not codeshare',a,r)
        same_utc_departures.append(r)
        active=[old for old in active if old['_arr_utc'] and old['_arr_utc']>=r['_dep_utc']]
        for a in active:
            if (a['flight_number']==r['flight_number'] and
                a['departure_date']==r['departure_date']):continue
            if not different_destination_cities(a,r,airports):continue
            ids=tuple(sorted((id(a),id(r))))
            if ids in overlap_pairs:continue
            overlap_pairs.add(ids)
            mark('flight_overlap',
                f'Second departure starts before first arrival (UTC): '+
                f'{a["_arr_utc"].isoformat()} > {r["_dep_utc"].isoformat()}',a,r)
        if r['_arr_utc'] is not None:active.append(r)

        if i==0:continue
        prev=timed[i-1]
        minutes=(r['_dep_utc']-prev['_dep_utc']).total_seconds()/60
        distance=haversine_km((prev['origin_airport'] or '').upper(),
                              (r['origin_airport'] or '').upper(),airports)
        if minutes>0 and distance is not None and distance>=args.min_speed_distance_km:
            speed=distance/(minutes/60)
            if speed>args.max_speed_kmh:
                mark('impossible_speed',
                    f'Departure origins {distance:.0f} km apart in {minutes:.0f} min '+
                    f'({speed:.0f} km/h > {args.max_speed_kmh:.0f})',prev,r)
        if (prev['_arr_utc'] and prev['destination_airport'] and r['origin_airport'] and
            prev['destination_airport'].upper()!=r['origin_airport'].upper()):
            ground_minutes=(r['_dep_utc']-prev['_arr_utc']).total_seconds()/60
            if 0<=ground_minutes<args.airport_change_minutes:
                mark('short_airport_change',
                    f'{ground_minutes:.0f} min between arrival at '+
                    f'{prev["destination_airport"]} and next departure at '+
                    f'{r["origin_airport"]}',prev,r)
        elif (not prev['_arr_utc'] and prev['origin_airport'] and r['origin_airport']
              and prev['origin_airport'].upper()!=r['origin_airport'].upper()
              and 0<=minutes<args.short_departure_minutes):
            mark('short_airport_change',
                f'{minutes:.0f} min between departures from different airports; '+
                'arrival time of the first flight is unknown',prev,r)

    points={k:(WEIGHTS[k] if count[k]>0 else 0) for k in WEIGHTS}
    category_points={label:sum(points[k] for k in indicators)
                     for label,indicators in CATEGORIES.items()}
    raw=sum(points.values());score=min(10,raw)
    causes=[k for k in WEIGHTS if points[k]>0]
    stats['persons']+=1
    stats['flights']+=len(flights)
    stats['foreign_flights']+=foreign
    stats['country_comparable']+=country_comparable
    stats['full_utc_intervals']+=full_intervals
    for k in WEIGHTS:
        if count[k]:stats['persons_with_'+k]+=1
    row={
        'rank':'','person_id':pid,'passenger_name':display_name,
        'passport':'|'.join(sorted(docs)),
        'score_10':score,'raw_points':raw,'is_suspect':int(score>5),
        'categories_ru':' | '.join(c for c in CATEGORIES if category_points[c]),
        'false_documents_points':category_points['Ложные документы'],
        'false_flight_points':category_points['Ложная запись о полёте'],
        'unusual_behavior_points':category_points['Нетипичное поведение'],
        'score_breakdown':'; '.join(f'{DESCRIPTIONS[k]} +{points[k]}' for k in causes),
        'reasons_ru':' | '.join(DESCRIPTIONS[k] for k in causes),
        'flights':len(flights),'max_flights_in_year':max(years.values(),default=0),
        'foreign_flights':foreign,'flight_intervals_utc':full_intervals,
    }
    row.update({k:count[k] for k in WEIGHTS})
    return row


def write_excel(filename, candidates, stats,meta_source):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RuntimeError('Install Excel support: py -3.12 -m pip install openpyxl') from e
    wb=Workbook();ws=wb.active;ws.title='Кандидаты'
    columns=[('rank','Ранг',9),('person_id','Person ID',24),
        ('passenger_name','ФИО',30),('passport','Паспорт',21),
        ('score_10','Баллы / 10',13),('raw_points','Баллы без лимита',17),
        ('categories_ru','Категории',50),
        ('false_documents_points','Документы',15),
        ('false_flight_points','Запись о полёте',19),
        ('unusual_behavior_points','Поведение',15),
        ('flights','Кол-во рейсов',15),('max_flights_in_year','Макс. рейсов в год',18),
        ('foreign_flights','Заграничных рейсов',20),
        ('reasons_ru','Признаки',82),('score_breakdown','Начисление баллов',95)]
    ws.append([n for k,n,w in columns])
    for i,row in enumerate(candidates,2):
        vals=[row.get(k,'') for k,n,w in columns]
        for j,(key,_,_) in enumerate(columns):
            if key in ('passport','person_id'):vals[j]=str(vals[j])
        ws.append(vals)
        for cell in ws[i]:cell.alignment=Alignment(vertical='top',wrap_text=cell.column in (7,14,15))
        ws.row_dimensions[i].height=47
        ws.cell(i,5).font=Font(bold=True,color='1F3650')
        if i%2==0:
            for c in ws[i]:c.fill=PatternFill('solid',fgColor='F3F7FB')
    for i,(_,_,width) in enumerate(columns,1):
        ws.column_dimensions[get_column_letter(i)].width=width
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='1B3A55')
        cell.font=Font(bold=True,color='FFFFFF',size=11)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=34;ws.freeze_panes='E2';ws.auto_filter.ref=f'A1:O{max(2,ws.max_row)}'
    ws.sheet_view.showGridLines=False
    meth=wb.create_sheet('Методика');meth.sheet_view.showGridLines=False
    notes=[('Правило','Баллы 0..10: минимум 6 баллов для включения в список кандидатов'),
        ('Ограничение','Каждый критерий начисляет баллы один раз на person_id; повторные события лишь учитываются'),
        ('Исходная сумма','Поле raw_points может превышать 10; score_10 ограничено 10'),
        ('Время','Местное время сравнивается после перевода в UTC через часовые пояса аэропортов'),
        ('Имя','Транслитерация: '+TRANSLITERATOR),
        ('Аэропорты',meta_source),
        ('Проверено идентификаторов',stats['persons']),
        ('Оговорка','Это учебные индикаторы качества данных и маршрутов, НЕ вывод о шпионаже')]
    for r in notes:meth.append(r)
    meth.append([]);meth.append(['Категория','Признак','Баллы'])
    for cat,ks in CATEGORIES.items():
        for k in ks:meth.append([cat,DESCRIPTIONS[k],WEIGHTS[k]])
    meth.column_dimensions['A'].width=34
    meth.column_dimensions['B'].width=94
    meth.column_dimensions['C'].width=14
    for row in meth:
        row[0].font=Font(bold=True,color='1B3A55')
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    wb.save(filename)


def run(args):
    args.out.mkdir(parents=True,exist_ok=True)
    airport_info,airport_source=airport_metadata(args.airport_meta)
    stats=Counter();summaries=[]
    path=args.out/'04_anomaly_events.csv'
    with open_csv(args.cleaned) as f,open_csv(path,'w') as evt_f:
        rd=csv.DictReader(f,delimiter=csv_separator(f))
        required={'person_id','passenger_name','passport','date_of_birth',
            'flight_number','departure_date','departure_time','arrival_date','arrival_time',
            'origin_airport','destination_airport','origin_country','destination_country',
            'source','source_record_id'}
        if not required.issubset(rd.fieldnames or ()):
            raise ValueError('Input is not 03_cleaned.csv; missing '+str(sorted(required-set(rd.fieldnames or []))))
        evt=csv.DictWriter(evt_f,fieldnames=EVENT_FIELDS,delimiter=';');evt.writeheader()
        def write_event(row):
            stats['events_found']+=1
            if stats['events_written']<args.max_events:
                evt.writerow(row);stats['events_written']+=1
        prev_id=None
        for pid,iter_rows in groupby(rd,key=lambda r:r['person_id']):
            if prev_id is not None and pid<prev_id:
                raise ValueError('CSV must be sorted by person_id (use current stage 3 output)')
            result=analyze_person(pid,list(iter_rows),airport_info,args,write_event,stats)
            if result:summaries.append(result)
            prev_id=pid
            if stats['persons']%25000==0:
                print('Processed',f'{stats["persons"]:,}','passenger IDs',flush=True)
    summaries.sort(key=lambda r:(-r['score_10'],-r['raw_points'],
        -r['false_documents_points'],-r['false_flight_points'],r['person_id']))
    for n,r in enumerate(summaries,1):r['rank']=n
    cutoff=max(6,args.min_score)
    suspects=[r for r in summaries if r['score_10']>=cutoff]
    for file_name,contents in (
        ('04_passenger_summary.csv',summaries),
        ('04_all_ranked_candidates.csv',suspects),
        ('04_review_candidates.csv',suspects[:args.top])):
        with open_csv(args.out/file_name,'w') as f:
            w=csv.DictWriter(f,fieldnames=BASE_FIELDS,delimiter=';')
            w.writeheader();w.writerows(contents)
    report={
        'rules':'User-specified 3 categories; each indicator adds its weight at most once per person',
        'score':'score_10 = min(10, sum(weights for triggered indicators)); candidate iff score_10 > 5',
        'candidate_min_score':cutoff,
        'passenger_ids':len(summaries),'candidates':len(suspects),
        'top_preview':min(len(suspects),args.top),
        'score_weights':WEIGHTS,'categories':CATEGORIES,
        'transliteration':TRANSLITERATOR,
        'airport_metadata':airport_source,
        'coverage':{'utc_full_flight_intervals':stats['full_utc_intervals'],
                    'crossborder_comparable_flights':stats['country_comparable'],
                    'detected_crossborder_flights':stats['foreign_flights'],
                    'events_found':stats['events_found'],
                    'events_written':stats['events_written']},
        'assumptions':[
            'An identity is person_id assigned at stage 3; a passport is not necessarily verified.',
            'Only non-empty and sufficiently complete full names are compared.',
            'Same-date same-number route variants are checked for discrepancies, not independently timed as two trips.',
            'Code-shared departures to the same city are not counted as a flight overlap.',
            'Flight time comparisons require correct timezones; without an arrival, interval overlaps are not guessed.',
            'Impossible speed is approximate departure-origin distance divided by UTC time between departures.',
            'Airport transfer under 180 min is a heuristic and may be feasible at nearby airports.',
            'Foreign flight means at least one flight crossing known national borders; max +3 per person.',
            'A score is not probability of espionage, and source data does not prove boarding.',
        ],
        'stats':dict(stats),
    }
    (args.out/'04_analysis_report.json').write_text(
        json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    if args.excel:
        write_excel(args.out/'04_review_candidates.xlsx',suspects[:args.top],stats,airport_source)
    else:
        (args.out/'04_review_candidates.xlsx').unlink(missing_ok=True)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('cleaned',type=Path,help='03_cleaned.csv sorted by person_id')
    p.add_argument('--out',type=Path,default=Path('result/04_analysis'))
    p.add_argument('--airport-meta',type=Path,help='Optional CSV: iata;country;timezone;lat;lon;city')
    p.add_argument('--max-speed-kmh',type=float,default=1200)
    p.add_argument('--min-speed-distance-km',type=float,default=200)
    p.add_argument('--airport-change-minutes',type=int,default=180)
    p.add_argument('--short-departure-minutes',type=int,default=90)
    p.add_argument('--min-score',type=int,default=6,
                   help='Min score for candidates, but never less than 6')
    p.add_argument('--top',type=int,default=300)
    p.add_argument('--max-events',type=int,default=100000)
    p.add_argument('--excel',action='store_true')
    args=p.parse_args()
    if args.min_score>10:p.error('--min-score must be 0..10')
    if args.top<1:p.error('--top must be >=1')
    if args.max_speed_kmh<=0:p.error('--max-speed-kmh must be positive')
    run(args)


if __name__=='__main__':
    main()
